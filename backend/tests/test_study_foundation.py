from __future__ import annotations

import io
import uuid

import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models  # noqa: F401 - register all SQLAlchemy models
from app.main import app
from core.config import settings
from database.base import Base
from database.session import get_db
from models.study import (
    StudyAttempt,
    StudyCard,
    StudyDocument,
    StudyDocumentUnit,
    StudyQuiz,
    StudyUserCardState,
)
from models.user import User
from services.study_files import LocalStudyStorage
from services.study_ocr import OCRBlock
from services.study_parser import parse_and_store_document
from utils.auth import create_token, hash_password


@pytest.fixture()
def study_api(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "study_upload_dir", str(tmp_path / "study_uploads"))
    # Unit tests inject a deterministic OCR engine instead of loading model weights.
    monkeypatch.setattr(settings, "study_ocr_enabled", False)
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_factory = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        with session_factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    try:
        yield TestClient(app), session_factory
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def _user(session_factory, label: str) -> tuple[str, dict[str, str]]:
    with session_factory() as db:
        user = User(
            student_id=f"{label}-{uuid.uuid4().hex[:8]}",
            name=label,
            nickname=label,
            password_hash=hash_password("secure-pass"),
        )
        db.add(user)
        db.commit()
        db.refresh(user)
        user_id = user.id
    return user_id, {"Authorization": f"Bearer {create_token(user_id)}"}


def _pptx_bytes(slides: tuple[tuple[str, str], ...] | None = None) -> bytes:
    slides = slides or (("Transport Layer", "TCP is connection-oriented and reliable."),)
    presentation = Presentation()
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _pdf_bytes(*pages: str) -> bytes:
    document = pymupdf.open()
    try:
        for text in pages:
            page = document.new_page()
            if text:
                page.insert_text((72, 72), text)
        return document.tobytes()
    finally:
        document.close()


def _document(session_factory, user_id: str, filename: str) -> str:
    with session_factory() as db:
        document = StudyDocument(
            user_id=user_id,
            original_filename=filename,
            file_type="pdf",
            storage_path=f"study/{user_id}/{filename}",
            file_size=128,
            sha256="a" * 64,
        )
        db.add(document)
        db.commit()
        db.refresh(document)
        return document.id


def test_study_documents_require_login_and_are_owner_scoped(study_api):
    client, session_factory = study_api
    owner_id, owner_headers = _user(session_factory, "study-owner")
    _, other_headers = _user(session_factory, "study-other")
    document_id = _document(session_factory, owner_id, "network.pdf")

    assert client.get("/api/v1/study/documents").status_code == 401

    owner_list = client.get("/api/v1/study/documents", headers=owner_headers)
    assert owner_list.status_code == 200, owner_list.text
    assert owner_list.json()["data"]["total"] == 1
    assert owner_list.json()["data"]["documents"][0]["id"] == document_id

    other_list = client.get("/api/v1/study/documents", headers=other_headers)
    assert other_list.status_code == 200, other_list.text
    assert other_list.json()["data"] == {"total": 0, "documents": []}

    owner_detail = client.get(
        f"/api/v1/study/documents/{document_id}", headers=owner_headers
    )
    assert owner_detail.status_code == 200, owner_detail.text
    assert owner_detail.json()["data"]["original_filename"] == "network.pdf"

    hidden_detail = client.get(
        f"/api/v1/study/documents/{document_id}", headers=other_headers
    )
    assert hidden_detail.status_code == 404, hidden_detail.text


@pytest.mark.parametrize(
    ("filename", "content", "content_type", "file_type"),
    [
        ("network.pdf", b"%PDF-1.7\nminimal-test-pdf", "application/pdf", "pdf"),
        (
            "lecture.pptx",
            _pptx_bytes(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "pptx",
        ),
    ],
    ids=["pdf", "pptx"],
)
def test_authenticated_user_can_upload_supported_study_file(
    study_api, filename, content, content_type, file_type
):
    client, session_factory = study_api
    user_id, headers = _user(session_factory, f"upload-{file_type}")

    response = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={"file": (filename, content, content_type)},
    )

    assert response.status_code == 201, response.text
    payload = response.json()["data"]
    assert payload["original_filename"] == filename
    assert payload["file_type"] == file_type
    assert payload["file_size"] == len(content)
    assert payload["status"] == "uploaded"

    with session_factory() as db:
        document = db.get(StudyDocument, payload["id"])
        assert document is not None
        assert document.user_id == user_id
        assert not document.storage_path.startswith(("/", "\\"))
        stored_path = LocalStudyStorage().resolve(document.storage_path)
        assert stored_path.read_bytes() == content


def test_study_upload_requires_login_and_rejects_invalid_files(study_api):
    client, session_factory = study_api
    _, headers = _user(session_factory, "invalid-upload")

    unauthenticated = client.post(
        "/api/v1/study/documents",
        files={"file": ("notes.pdf", b"%PDF-1.7\ncontent", "application/pdf")},
    )
    assert unauthenticated.status_code == 401

    wrong_extension = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )
    assert wrong_extension.status_code == 422
    assert wrong_extension.json()["message"] == "仅支持 PDF 或 PPTX 文件"

    fake_pdf = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={"file": ("notes.pdf", b"not a pdf", "application/pdf")},
    )
    assert fake_pdf.status_code == 422
    assert fake_pdf.json()["message"] == "文件内容不是有效的 PDF 格式"

    fake_pptx = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={"file": ("slides.pptx", b"PK-not-office", "application/zip")},
    )
    assert fake_pptx.status_code == 422
    assert fake_pptx.json()["message"] == "文件内容不是有效的 PPTX 格式"


def test_oversized_study_upload_is_removed_without_database_record(
    study_api, monkeypatch
):
    client, session_factory = study_api
    _, headers = _user(session_factory, "oversized-upload")
    monkeypatch.setattr(settings, "study_upload_max_bytes", 1024 * 1024)
    content = b"%PDF-1.7\n" + b"x" * (1024 * 1024)

    response = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={"file": ("large.pdf", content, "application/pdf")},
    )

    assert response.status_code == 422
    assert response.json()["message"] == "学习资料不能超过 1 MB"
    with session_factory() as db:
        assert db.query(StudyDocument).count() == 0
    storage = LocalStudyStorage()
    assert not list(storage.base_dir.rglob("*.uploading"))


def test_pdf_parse_creates_page_units_and_is_idempotent(study_api):
    client, session_factory = study_api
    _, owner_headers = _user(session_factory, "pdf-owner")
    _, other_headers = _user(session_factory, "pdf-other")
    content = _pdf_bytes(
        "TCP is connection-oriented and reliable.",
        "UDP is connectionless.",
    )
    uploaded = client.post(
        "/api/v1/study/documents",
        headers=owner_headers,
        files={"file": ("transport.pdf", content, "application/pdf")},
    )
    assert uploaded.status_code == 201, uploaded.text
    document_id = uploaded.json()["data"]["id"]

    hidden = client.post(
        f"/api/v1/study/documents/{document_id}/parse", headers=other_headers
    )
    assert hidden.status_code == 404

    parsed = client.post(
        f"/api/v1/study/documents/{document_id}/parse", headers=owner_headers
    )
    assert parsed.status_code == 200, parsed.text
    assert parsed.json()["data"]["document"]["status"] == "parsed"
    assert parsed.json()["data"]["document"]["page_count"] == 2
    assert parsed.json()["data"]["unit_count"] == 2

    units_response = client.get(
        f"/api/v1/study/documents/{document_id}/units", headers=owner_headers
    )
    assert units_response.status_code == 200, units_response.text
    units = units_response.json()["data"]["units"]
    assert [unit["page_number"] for unit in units] == [1, 2]
    assert [unit["unit_index"] for unit in units] == [0, 0]
    assert "TCP is connection-oriented" in units[0]["normalized_text"]
    assert "UDP is connectionless" in units[1]["normalized_text"]
    assert client.get(
        f"/api/v1/study/documents/{document_id}/units", headers=other_headers
    ).status_code == 404

    repeated = client.post(
        f"/api/v1/study/documents/{document_id}/parse", headers=owner_headers
    )
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["data"]["unit_count"] == 2
    with session_factory() as db:
        assert db.query(StudyDocumentUnit).filter_by(document_id=document_id).count() == 2


def test_pptx_parse_creates_one_unit_per_nonempty_slide(study_api):
    client, session_factory = study_api
    _, headers = _user(session_factory, "pptx-owner")
    content = _pptx_bytes((
        ("Transport Layer", "TCP is reliable."),
        ("Network Layer", "IP provides addressing."),
    ))
    uploaded = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={
            "file": (
                "network.pptx",
                content,
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    document_id = uploaded.json()["data"]["id"]

    parsed = client.post(
        f"/api/v1/study/documents/{document_id}/parse", headers=headers
    )
    assert parsed.status_code == 200, parsed.text
    assert parsed.json()["data"]["document"]["page_count"] == 2

    units_response = client.get(
        f"/api/v1/study/documents/{document_id}/units", headers=headers
    )
    units = units_response.json()["data"]["units"]
    assert [unit["unit_type"] for unit in units] == ["slide", "slide"]
    assert "Transport Layer" in units[0]["raw_text"]
    assert "TCP is reliable" in units[0]["raw_text"]


def test_parse_records_explicit_failure_for_document_without_text(study_api):
    client, session_factory = study_api
    _, headers = _user(session_factory, "blank-pdf")
    uploaded = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={"file": ("blank.pdf", _pdf_bytes(""), "application/pdf")},
    )
    document_id = uploaded.json()["data"]["id"]

    parsed = client.post(
        f"/api/v1/study/documents/{document_id}/parse", headers=headers
    )
    assert parsed.status_code == 422
    assert parsed.json()["message"] == "PDF 中没有可提取文字，且 OCR 未启用"

    detail = client.get(
        f"/api/v1/study/documents/{document_id}", headers=headers
    )
    assert detail.json()["data"]["status"] == "failed"
    assert detail.json()["data"]["error_code"] == "no_extractable_text"


def test_image_pdf_uses_injected_ocr_and_persists_blocks(study_api):
    client, session_factory = study_api
    user_id, headers = _user(session_factory, "ocr-pdf")
    uploaded = client.post(
        "/api/v1/study/documents",
        headers=headers,
        files={"file": ("image.pdf", _pdf_bytes(""), "application/pdf")},
    )
    document_id = uploaded.json()["data"]["id"]

    class FakeOCREngine:
        def recognize(self, pixmap):
            assert pixmap.width > 0 and pixmap.height > 0
            return [
                OCRBlock(
                    "土的三相组成",
                    0.998,
                    [
                        int(pixmap.width * 0.1),
                        int(pixmap.height * 0.1),
                        int(pixmap.width * 0.5),
                        int(pixmap.height * 0.2),
                    ],
                ),
                OCRBlock(
                    "固相、液相、气相",
                    0.962,
                    [
                        int(pixmap.width * 0.1),
                        int(pixmap.height * 0.3),
                        int(pixmap.width * 0.6),
                        int(pixmap.height * 0.4),
                    ],
                ),
            ]

    with session_factory() as db:
        document = db.get(StudyDocument, document_id)
        assert document is not None
        unit_count = parse_and_store_document(
            db,
            document,
            ocr_engine=FakeOCREngine(),
        )
        assert unit_count == 1

    units_response = client.get(
        f"/api/v1/study/documents/{document_id}/units", headers=headers
    )
    assert units_response.status_code == 200, units_response.text
    unit = units_response.json()["data"]["units"][0]
    assert unit["extraction_method"] == "ocr"
    assert unit["ocr_confidence"] == pytest.approx(0.98)
    assert unit["raw_text"] == "土的三相组成\n固相、液相、气相"
    assert unit["safe_text"] == "土的三相组成\n\n固相、液相、气相"
    assert unit["quality_status"] == "accepted"
    assert unit["quality_reasons"] == []
    assert len(unit["ocr_blocks"]) == 2
    assert all(block["bbox_unit"] == "relative" for block in unit["ocr_blocks"])
    assert all(block["decision"] == "accepted" for block in unit["ocr_blocks"])


def test_study_models_persist_the_minimal_learning_chain(study_api):
    _, session_factory = study_api
    user_id, _ = _user(session_factory, "study-chain")
    document_id = _document(session_factory, user_id, "tcp.pdf")

    with session_factory() as db:
        unit = StudyDocumentUnit(
            document_id=document_id,
            page_number=17,
            unit_index=0,
            unit_type="paragraph",
            raw_text="TCP是一种面向连接、可靠的传输层协议。",
            normalized_text="TCP是一种面向连接、可靠的传输层协议。",
            text_hash="b" * 64,
        )
        db.add(unit)
        db.flush()

        card = StudyCard(
            document_id=document_id,
            source_unit_id=unit.id,
            front="TCP 是一种什么协议？",
            back="TCP 是一种面向连接、可靠的传输层协议。",
            highlights=["面向连接", "可靠"],
            sources=[{
                "page": 17,
                "quote": "TCP是一种面向连接、可靠的传输层协议。",
            }],
            status="validated",
        )
        db.add(card)
        db.flush()

        quiz = StudyQuiz(
            card_id=card.id,
            quiz_type="fill_blank",
            prompt="TCP 是一种____、可靠的传输层协议。",
            options=None,
            answer="面向连接",
            explanation="原文指出 TCP 面向连接且可靠。",
        )
        db.add(quiz)
        db.flush()

        db.add(StudyAttempt(
            quiz_id=quiz.id,
            user_id=user_id,
            user_answer="面向连接",
            is_correct=True,
        ))
        db.add(StudyUserCardState(
            user_id=user_id,
            card_id=card.id,
            status="learning",
            study_count=1,
            correct_count=1,
        ))
        db.commit()

        assert db.query(StudyDocumentUnit).count() == 1
        assert db.query(StudyCard).filter_by(status="validated").count() == 1
        assert db.query(StudyQuiz).count() == 1
        assert db.query(StudyAttempt).filter_by(is_correct=True).count() == 1
        state = db.query(StudyUserCardState).one()
        assert (state.study_count, state.correct_count, state.status) == (1, 1, "learning")


def test_parse_rebuilds_legacy_ocr_quality_but_caches_paragraph_gated_results(study_api):
    client, session_factory = study_api
    _, headers = _user(session_factory, "legacy-ocr")
    uploaded = client.post(
        "/api/v1/study/documents", headers=headers,
        files={"file": ("old.pdf", _pdf_bytes(""), "application/pdf")},
    )
    document_id = uploaded.json()["data"]["id"]

    class FakeOCREngine:
        calls = 0

        def recognize(self, pixmap):
            self.calls += 1
            width, height = pixmap.width, pixmap.height
            return [
                OCRBlock("同段第一行", 0.99, [int(width*.1), int(height*.1), int(width*.8), int(height*.14)]),
                OCRBlock("不可靠的最后一行", 0.5, [int(width*.1), int(height*.15), int(width*.8), int(height*.19)]),
            ]

    engine = FakeOCREngine()
    with session_factory() as db:
        document = db.get(StudyDocument, document_id)
        document.status = "parsed"
        db.add(StudyDocumentUnit(
            document_id=document_id, page_number=1, unit_index=0, unit_type="page",
            raw_text="旧的残句", normalized_text="旧的残句", safe_text="旧的残句",
            text_hash="c" * 64, extraction_method="ocr",
            ocr_blocks=[{"text": "旧的残句", "decision": "accepted"}],
        ))
        db.commit()
        assert parse_and_store_document(db, document, ocr_engine=engine) == 1
        assert engine.calls == 2
        unit = db.query(StudyDocumentUnit).filter_by(document_id=document_id).one()
        assert unit.safe_text == "" and unit.quality_status == "rejected"
        assert all(block["paragraph_gate_version"] == "paragraph-v1" for block in unit.ocr_blocks)
        assert parse_and_store_document(db, document, ocr_engine=engine) == 1
        assert engine.calls == 2
