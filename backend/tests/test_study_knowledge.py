from __future__ import annotations

import hashlib
import json
import uuid
from pathlib import Path
from types import SimpleNamespace

import httpx
import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models  # noqa: F401
from app.main import app
from database.base import Base
from database.session import get_db
from models.study import (
    StudyDocument,
    StudyDocumentUnit,
    StudyKnowledgeRun,
    StudyKnowledgePageCheckpoint,
    StudyKnowledgeUnit,
    StudyStructuredBlock,
)
from models.user import User
from services.study_knowledge import (
    ModelResult,
    StructuredBlock,
    StructuredDocument,
    StructuredSourceSpan,
    create_or_reuse_run,
    extract_from_structured_document,
    process_knowledge_run,
)
from services.study_knowledge_v2 import ModelResult as ModelResultV2
from services.study_knowledge_v2 import extract_from_numbered_document
from services.study_knowledge_doubao import (
    DirectModelResult,
    DoubaoFilesResponsesModel,
    DoubaoPageImagesResponsesModel,
    _filter_reasons,
    _source_as_pdf,
    DirectKnowledgePoint,
    extract_direct_document,
)
from services.study_source_units import structured_from_pptx
from utils.auth import create_token, hash_password


class FakeKnowledgeModel:
    model = "fake-grounded-model"

    def extract(self, blocks):
        return ModelResult(payload={"candidates": [
            {
                "candidate_id": "definition-ok",
                "type": "definition",
                "block_id": "p1",
                "extraction_text": "市场细分是指根据消费者需求差异划分市场。",
                "risk_flags": ["none"],
            },
            {
                "candidate_id": "rewritten",
                "type": "definition",
                "block_id": "p1",
                "extraction_text": "市场细分是按照需求划分市场。",
                "risk_flags": ["none"],
            },
            {
                "candidate_id": "formula",
                "type": "definition",
                "block_id": "f1",
                "extraction_text": "孔隙比是e=Vv/Vs。",
                "risk_flags": ["formula_dependency"],
            },
            {
                "candidate_id": "list-fragment",
                "type": "complete_list",
                "block_id": "p2",
                "extraction_text": "战略具有全局性、长远性。",
                "risk_flags": ["possible_incomplete"],
            },
            {
                "candidate_id": "list-ok",
                "type": "complete_list",
                "block_id": "p2",
                "extraction_text": "战略具有全局性、长远性、竞合性、纲领性、相对稳定性。",
                "risk_flags": ["none"],
            },
            {
                "candidate_id": "corrupt-list",
                "type": "complete_list",
                "block_id": "p3",
                "extraction_text": "按来源不同（中，中两方；中，中，中'合作；中外合作）",
                "risk_flags": ["none"],
            },
            {
                "candidate_id": "overbroad-definition",
                "type": "definition",
                "block_id": "p4",
                "extraction_text": "缓存是指暂时保存数据的机制。作用是减少重复读取。",
                "risk_flags": ["none"],
            },
        ]}, model=self.model, duration_ms=2.0)

    def review(self, blocks, candidates):
        return ModelResult(payload={"decisions": [
            {
                "candidate_id": "definition-ok",
                "decision": "usable",
                "error_types": ["none"],
                "related_source_block_ids": ["p1"],
            },
            {
                "candidate_id": "list-ok",
                "decision": "usable",
                "error_types": ["none"],
                "related_source_block_ids": ["p2"],
            }
        ]}, model=self.model, duration_ms=3.0)


def _structured_document() -> StructuredDocument:
    return StructuredDocument(
        revision_id="revision-1",
        source_sha256="a" * 64,
        parser_name="test-parser",
        parser_version="1",
        parser_config={},
        blocks=[
            StructuredBlock(
                id="h1", type="heading", raw_text="市场战略", reading_order=0,
                source_spans=[StructuredSourceSpan(1, [0.1, 0.1, 0.5, 0.2], 0, 4)],
                quality_signals={"usable_text": True},
            ),
            StructuredBlock(
                id="p1", type="paragraph",
                raw_text="市场细分是指根据消费者需求差异划分市场。", reading_order=1,
                source_spans=[StructuredSourceSpan(1, [0.1, 0.2, 0.9, 0.3], 0, 20)],
                quality_signals={"usable_text": True},
            ),
            StructuredBlock(
                id="p2", type="paragraph",
                raw_text="战略具有全局性、长远性、竞合性、纲领性、相对稳定性。", reading_order=2,
                source_spans=[StructuredSourceSpan(1, [0.1, 0.3, 0.9, 0.4], 0, 26)],
                quality_signals={"usable_text": True},
            ),
            StructuredBlock(
                id="f1", type="formula", raw_text="孔隙比是e=Vv/Vs。", reading_order=3,
                source_spans=[StructuredSourceSpan(1, [0.1, 0.4, 0.9, 0.5], 0, 12)],
                quality_signals={"usable_text": False},
            ),
            StructuredBlock(
                id="p3", type="paragraph",
                raw_text="按来源不同（中，中两方；中，中，中'合作；中外合作）", reading_order=4,
                source_spans=[StructuredSourceSpan(1, [0.1, 0.5, 0.9, 0.6], 0, 27)],
                quality_signals={"usable_text": True},
            ),
            StructuredBlock(
                id="p4", type="paragraph",
                raw_text="缓存是指暂时保存数据的机制。作用是减少重复读取。", reading_order=5,
                source_spans=[StructuredSourceSpan(1, [0.1, 0.6, 0.9, 0.7], 0, 24)],
                quality_signals={"usable_text": True},
            ),
        ],
    )


def test_pipeline_rebuilds_exact_source_and_rejects_rewrite_formula_and_fragment():
    records, audit = extract_from_structured_document(
        _structured_document(), FakeKnowledgeModel(), max_chars=5000
    )
    by_id = {record["extraction_meta"]["candidate"]["candidate_id"]: record for record in records}

    accepted = by_id["definition-ok"]
    assert accepted["disposition"] == "usable"
    assert accepted["content"] == "市场细分是指根据消费者需求差异划分市场。"
    assert accepted["check_meta"]["exact_slice"] is True
    assert accepted["source_refs"][0]["block_id"] == "p1"
    assert accepted["source_refs"][0]["char_start"] == 0

    assert by_id["rewritten"]["disposition"] == "uncertain"
    assert by_id["rewritten"]["content"] == ""
    assert "source_slice_mismatch" in by_id["rewritten"]["reasons"]
    assert by_id["formula"]["disposition"] == "unsupported"
    assert "unsupported_block_type:formula" in by_id["formula"]["reasons"]
    assert by_id["list-fragment"]["disposition"] == "uncertain"
    assert "model_risk:possible_incomplete" in by_id["list-fragment"]["reasons"]
    assert by_id["list-ok"]["disposition"] == "usable"
    assert by_id["list-ok"]["knowledge_type"] == "complete_list"
    assert by_id["corrupt-list"]["disposition"] == "uncertain"
    assert "suspicious_mixed_script_glyph" in by_id["corrupt-list"]["reasons"]
    assert by_id["overbroad-definition"]["disposition"] == "uncertain"
    assert "definition_shape_incomplete" in by_id["overbroad-definition"]["reasons"]
    assert audit["model_call_count"] == 2


def _api_fixture():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    factory = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app), factory, engine


def _create_user(factory, label):
    with factory() as db:
        user = User(
            student_id=f"{label}-{uuid.uuid4().hex[:8]}",
            name=label,
            nickname=label,
            password_hash=hash_password("pass"),
        )
        db.add(user)
        db.commit()
        db.refresh(user)
        return user.id, {"Authorization": f"Bearer {create_token(user.id)}"}


def test_knowledge_api_is_owner_scoped_and_reuses_same_revision(monkeypatch):
    client, factory, engine = _api_fixture()
    monkeypatch.setattr(
        "app.api.v1.study.process_knowledge_run_by_id", lambda _run_id: None
    )
    try:
        owner_id, owner_headers = _create_user(factory, "owner")
        _, other_headers = _create_user(factory, "other")
        with factory() as db:
            document = StudyDocument(
                user_id=owner_id,
                original_filename="notes.pdf",
                file_type="pdf",
                storage_path="private/notes.pdf",
                file_size=10,
                sha256="b" * 64,
                status="parsed",
            )
            db.add(document)
            db.commit()
            db.refresh(document)
            document_id = document.id

        first = client.post(
            f"/api/v1/study/documents/{document_id}/knowledge-runs",
            headers=owner_headers,
            json={"force": False},
        )
        assert first.status_code == 202, first.text
        first_id = first.json()["data"]["run"]["id"]
        assert first.json()["data"]["reused"] is False
        assert first.json()["data"]["run"]["pipeline_version"] == "doubao-vision-v2"

        repeated = client.post(
            f"/api/v1/study/documents/{document_id}/knowledge-runs",
            headers=owner_headers,
            json={"force": False},
        )
        assert repeated.json()["data"]["run"]["id"] == first_id
        assert repeated.json()["data"]["reused"] is True

        forced = client.post(
            f"/api/v1/study/documents/{document_id}/knowledge-runs",
            headers=owner_headers,
            json={"force": True},
        )
        assert forced.json()["data"]["run"]["revision"] == 2
        assert forced.json()["data"]["run"]["id"] != first_id

        v2 = client.post(
            f"/api/v1/study/documents/{document_id}/knowledge-runs",
            headers=owner_headers,
            json={"force": False, "pipeline_version": "knowledge-v2"},
        )
        assert v2.status_code == 202, v2.text
        assert v2.json()["data"]["run"]["revision"] == 3
        assert v2.json()["data"]["run"]["pipeline_version"] == "knowledge-v2"

        hidden = client.get(
            f"/api/v1/study/knowledge-runs/{first_id}", headers=other_headers
        )
        assert hidden.status_code == 404
        own = client.get(
            f"/api/v1/study/knowledge-runs/{first_id}", headers=owner_headers
        )
        assert own.status_code == 200
        hidden_units = client.get(
            f"/api/v1/study/knowledge-runs/{first_id}/knowledge-units",
            headers=other_headers,
        )
        assert hidden_units.status_code == 404

        with factory() as db:
            document = db.get(StudyDocument, document_id)
            same, reused = create_or_reuse_run(db, document, force=False)
            assert reused is True
            assert same.revision == 1
            assert db.query(StudyKnowledgeRun).count() == 3
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


class RuntimeFakeModel:
    model = "runtime-fake"

    def extract(self, blocks):
        block = next(item for item in blocks if "缓存是指" in item["text"])
        return ModelResult(payload={"candidates": [{
            "candidate_id": "runtime-definition",
            "type": "definition",
            "source_ids": [block["source_id"]],
            "cloze_answers": ["缓存"],
            "risk_flags": ["none"],
        }]}, model=self.model, duration_ms=1.0)

    def review(self, candidates):
        return ModelResult(payload={"decisions": [{
            "candidate_id": "runtime-definition",
            "decision": "usable",
            "error_types": ["none"],
            "related_source_block_ids": candidates[0]["source_ids"],
        }]}, model=self.model, duration_ms=1.0)


class RuleReviewModel:
    model = "rule-review-fake"

    def extract(self, blocks):
        return ModelResult(payload={"candidates": []}, model=self.model, duration_ms=1.0)

    def review(self, blocks, candidates):
        return ModelResult(payload={"decisions": [
            {
                "candidate_id": item["candidate_id"],
                "decision": "usable",
                "error_types": ["none"],
                "related_source_block_ids": [item["block_id"]],
            }
            for item in candidates
        ]}, model=self.model, duration_ms=1.0)


def test_explicit_characteristic_list_rule_recovers_exact_subrange_for_review():
    text = "前文。\n特点：\n1）延迟低；\n2）吞吐高；\n3）可扩展。\n后文。"
    document = StructuredDocument(
        revision_id="rule-revision",
        source_sha256="e" * 64,
        parser_name="test",
        parser_version="1",
        parser_config={},
        blocks=[StructuredBlock(
            id="rule-block",
            type="paragraph",
            raw_text=text,
            reading_order=0,
            source_spans=[StructuredSourceSpan(1, None, 0, len(text))],
            quality_signals={"usable_text": True},
        )],
    )
    records, audit = extract_from_structured_document(document, RuleReviewModel())
    assert len(records) == 1
    assert records[0]["content"] == "特点：\n1）延迟低；\n2）吞吐高；\n3）可扩展。"
    assert records[0]["disposition"] == "usable"
    assert records[0]["extraction_meta"]["origin"] == "explicit_labeled_list_rule"
    assert records[0]["check_meta"]["exact_slice"] is True
    assert audit["model_call_count"] == 2


def test_database_run_persists_structured_revision_and_grounded_unit():
    _, factory, engine = _api_fixture()
    try:
        user_id, _ = _create_user(factory, "runtime")
        with factory() as db:
            document = StudyDocument(
                user_id=user_id,
                original_filename="runtime.pdf",
                file_type="pdf",
                storage_path="private/runtime.pdf",
                file_size=10,
                sha256="c" * 64,
                status="parsed",
            )
            db.add(document)
            db.flush()
            db.add(StudyDocumentUnit(
                document_id=document.id,
                page_number=1,
                unit_index=0,
                unit_type="page",
                raw_text="缓存是指暂时保存数据以减少重复读取的机制。",
                normalized_text="缓存是指暂时保存数据以减少重复读取的机制。",
                safe_text="缓存是指暂时保存数据以减少重复读取的机制。",
                text_hash="d" * 64,
                extraction_method="native",
                quality_status="accepted",
            ))
            db.commit()
            db.refresh(document)
            run, reused = create_or_reuse_run(
                db,
                document,
                pipeline_version="knowledge-v2",
            )
            assert reused is False
            completed = process_knowledge_run(db, run, model=RuntimeFakeModel())
            assert completed.status == "completed", (
                completed.error_code,
                completed.error_message,
            )
            assert completed.structured_revision_id
            assert completed.statistics["usable_count"] == 1
            stored = db.query(StudyKnowledgeUnit).filter_by(run_id=run.id).one()
            assert stored.content == "缓存是指暂时保存数据以减少重复读取的机制。"
            assert stored.structured_revision_id == completed.structured_revision_id
            assert stored.source_refs[0]["page_number"] == 1
            assert db.query(StudyStructuredBlock).filter_by(run_id=run.id).count() == 1
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


class SourceIdFakeModel:
    model = "source-id-fake"

    def __init__(self, source_ids, cloze_answers=None):
        self.source_ids = source_ids
        self.cloze_answers = cloze_answers or ["道路景观"]
        self.review_called = False

    def extract(self, source_units):
        return ModelResultV2(payload={"candidates": [{
            "candidate_id": "candidate-1",
            "type": "definition",
            "source_ids": self.source_ids,
            "cloze_answers": self.cloze_answers,
            "risk_flags": ["none"],
        }]}, model=self.model, duration_ms=1.0)

    def review(self, candidates):
        self.review_called = True
        return ModelResultV2(payload={"decisions": [{
            "candidate_id": "candidate-1",
            "decision": "usable",
            "error_types": ["none"],
            "related_source_block_ids": self.source_ids,
        }]}, model=self.model, duration_ms=1.0)


def _numbered_document():
    common = {
        "type": "paragraph",
        "extraction_method": "pptx_native",
        "quality_signals": {
            "usable_text": True,
            "source_container_id": "p2:shape3:para0",
        },
    }
    texts = [
        "道路景观指运动时看到的道路及周围空间。",
        "静止时看到的是道路与环境的三维空间。",
        "前者为动态，后者为静态。",
    ]
    blocks = [
        StructuredBlock(
            id=f"p2.shape3.para0.s{index}",
            raw_text=text,
            reading_order=index,
            source_spans=[StructuredSourceSpan(2, [0.1, 0.2, 0.8, 0.5], 0, len(text))],
            **common,
        )
        for index, text in enumerate(texts, 1)
    ]
    return StructuredDocument(
        revision_id="numbered-revision",
        source_sha256="f" * 64,
        parser_name="test-source-units",
        parser_version="2",
        parser_config={},
        blocks=blocks,
    )


def test_v2_model_selects_ids_and_code_rebuilds_multi_sentence_definition():
    document = _numbered_document()
    source_ids = [block.id for block in document.blocks]
    model = SourceIdFakeModel(source_ids)
    records, audit = extract_from_numbered_document(document, model)

    assert len(records) == 1
    assert records[0]["disposition"] == "usable"
    assert records[0]["content"] == "".join(block.raw_text for block in document.blocks)
    assert records[0]["check_meta"]["source_ids"] == source_ids
    assert records[0]["check_meta"]["content_rebuilt_from_source_ids"] is True
    assert records[0]["extraction_meta"]["cloze_spans"] == [
        {"text": "道路景观", "start": 0, "end": 4}
    ]
    assert audit["selection_mode"] == "stable_source_ids"
    assert model.review_called is True


def test_v2_rejects_non_contiguous_source_ids_before_review():
    document = _numbered_document()
    model = SourceIdFakeModel([document.blocks[0].id, document.blocks[2].id])
    records, _ = extract_from_numbered_document(document, model)

    assert records[0]["disposition"] == "uncertain"
    assert "non_contiguous_source_ids" in records[0]["reasons"]
    assert model.review_called is False


def test_v2_keeps_grounded_knowledge_when_cloze_hint_is_ambiguous():
    document = _numbered_document()
    model = SourceIdFakeModel(
        [block.id for block in document.blocks],
        cloze_answers=["道路"],
    )
    records, _ = extract_from_numbered_document(document, model)

    assert records[0]["disposition"] == "usable"
    assert records[0]["extraction_meta"]["cloze_status"] == "uncertain"
    assert records[0]["extraction_meta"]["cloze_reasons"] == [
        "cloze_answer_not_unique_in_source"
    ]


def test_pptx_source_ids_are_stable_and_keep_native_sentence_and_list_units(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "道路铺装景观"
    frame = slide.placeholders[1].text_frame
    frame.clear()
    frame.paragraphs[0].text = "道路景观指运动中看到的景象。静止时是三维空间景象。"
    item = frame.add_paragraph()
    item.text = "整体性铺装"
    item_properties = item._p.get_or_add_pPr()
    item_properties.append(parse_xml(f'<a:buChar {nsdecls("a")} char="•"/>'))
    path = tmp_path / "source-units.pptx"
    presentation.save(path)

    first = structured_from_pptx(path, source_sha256="1" * 64)
    second = structured_from_pptx(path, source_sha256="1" * 64)

    assert first.revision_id == second.revision_id
    assert [block.id for block in first.blocks] == [block.id for block in second.blocks]
    assert any(block.type == "heading" and block.raw_text == "道路铺装景观" for block in first.blocks)
    sentences = [block for block in first.blocks if block.type == "paragraph"]
    assert [block.raw_text for block in sentences] == [
        "道路景观指运动中看到的景象。",
        "静止时是三维空间景象。",
    ]
    assert any(block.type == "list_item" and block.raw_text == "整体性铺装" for block in first.blocks)
    assert all(block.source_spans[0].page_number == 1 for block in first.blocks)


class DirectFakeModel:
    model = "doubao-direct-fake"

    def __init__(self, knowledge_points, *, file_transport="files_api"):
        self.knowledge_points = knowledge_points
        self.file_transport = file_transport
        self.calls = 0

    def extract_file(self, path, *, filename, mime_type):
        self.calls += 1
        assert path.is_file()
        assert filename
        assert mime_type in {
            "application/pdf",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
        return DirectModelResult(
            payload={"knowledge_points": self.knowledge_points},
            model=self.model,
            duration_ms=3.0,
            remote_file_deleted=True,
            file_transport=self.file_transport,
            source_page_count=1 if self.file_transport == "page_images" else None,
        )


class CheckpointDirectFakeModel(DirectFakeModel):
    file_transport = "page_images"

    def __init__(self, *, fail_page_two_attempts=0):
        super().__init__([], file_transport="page_images")
        self.fail_page_two_attempts = fail_page_two_attempts
        self.page_calls = []

    def extract_page(self, path, page_number):
        self.page_calls.append(page_number)
        if page_number == 2 and self.fail_page_two_attempts:
            self.fail_page_two_attempts -= 1
            raise TimeoutError("page two timed out")
        return DirectModelResult(
            payload={"knowledge_points": [_direct_point(
                f"第{page_number}页的可恢复提取知识点。", page=page_number
            )]},
            model=self.model,
            duration_ms=float(page_number),
            remote_file_deleted=None,
            file_transport="page_images",
            model_call_count=1,
            model_durations_ms=[float(page_number)],
            source_page_count=3,
        )


def _direct_point(
    content,
    *,
    knowledge_type="statement",
    confidence=0.96,
    risk_flags=None,
    correction_level="typo",
    page=1,
):
    return {
        "content": content,
        "type": knowledge_type,
        "source_location": {"page": page, "slide": None},
        "confidence": confidence,
        "correction_level": correction_level,
        "risk_flags": risk_flags or [],
        "completeness": "complete",
    }


def test_doubao_direct_filter_accepts_safe_points_and_rejects_risky_content(tmp_path):
    path = tmp_path / "notes.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    source = StudyDocument(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        original_filename="notes.pdf",
        file_type="pdf",
        storage_path="notes.pdf",
        file_size=path.stat().st_size,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        status="uploaded",
    )
    safe = "市场细分是指根据消费者需求差异划分市场。"
    model = DirectFakeModel([
        _direct_point(safe),
        _direct_point("孔隙比是e=Vv/Vs。", risk_flags=["formula_dependency"]),
        _direct_point("例题：请计算三个方案的优先度。"),
        _direct_point("企业战略是企业发展的总体谋划。", confidence=0.70),
        _direct_point(safe, confidence=0.90),
    ])

    records, audit, structured = extract_direct_document(source, path, model)

    assert model.calls == 1
    assert [record["content"] for record in records] == [safe]
    assert records[0]["disposition"] == "usable"
    assert records[0]["source_refs"][0]["locator_type"] == "page"
    assert records[0]["extraction_meta"]["confidence"] == 0.96
    assert structured.blocks == []
    assert structured.parser_config["local_source_blocks_persisted"] is False
    assert audit["candidate_count"] == 5
    assert audit["accepted_count"] == 1
    assert audit["filtered_count"] == 4
    assert audit["filtered_by_reason"]["model_risk:formula_dependency"] == 1
    assert audit["filtered_by_reason"]["blocked_formula_marker"] == 1
    assert audit["filtered_by_reason"]["blocked_exercise_marker"] == 1
    assert audit["filtered_by_reason"]["confidence_below_threshold"] == 1
    assert audit["filtered_by_reason"]["duplicate_lower_confidence"] == 1
    assert audit["runtime_semantic_validation"] is False


def test_doubao_bad_candidate_schema_filters_only_that_candidate(tmp_path):
    path = tmp_path / "notes.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    source = StudyDocument(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        original_filename="notes.pdf",
        file_type="pdf",
        storage_path="notes.pdf",
        file_size=path.stat().st_size,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        status="uploaded",
    )
    bad = _direct_point("例题：求三个方案的优先度。")
    bad["type"] = "exercise_or_example"
    bad["completeness"] = "incomplete_context"
    model = DirectFakeModel([
        _direct_point("企业使命是企业经营目的和社会责任的规定。"),
        bad,
    ])

    records, audit, _ = extract_direct_document(source, path, model)

    assert len(records) == 1
    assert audit["candidate_count"] == 2
    assert audit["filtered_count"] == 1
    assert audit["filtered_by_reason"]["schema_invalid:type"] == 1
    assert audit["filtered_by_reason"]["schema_invalid:completeness"] == 1


@pytest.mark.parametrize("pipeline_version,file_transport", [
    ("doubao-direct-v1", "files_api"),
    ("doubao-vision-v2", "page_images"),
])
def test_doubao_direct_run_does_not_require_parsed_document_units(
    tmp_path, monkeypatch, pipeline_version, file_transport
):
    monkeypatch.setattr("core.config.settings.study_upload_dir", str(tmp_path))
    _, factory, engine = _api_fixture()
    try:
        with factory() as db:
            user = User(
                student_id=f"direct-{uuid.uuid4().hex[:8]}",
                name="direct",
                nickname="direct",
                password_hash=hash_password("pass"),
            )
            db.add(user)
            db.flush()
            relative = f"{user.id}/source.pdf"
            path = tmp_path / relative
            path.parent.mkdir(parents=True)
            path.write_bytes(b"%PDF-1.4\n")
            document = StudyDocument(
                user_id=user.id,
                original_filename="source.pdf",
                file_type="pdf",
                storage_path=relative,
                file_size=path.stat().st_size,
                sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                status="uploaded",
            )
            db.add(document)
            db.commit()
            db.refresh(document)
            run, reused = create_or_reuse_run(
                db, document, pipeline_version=pipeline_version
            )
            assert reused is False
            completed = process_knowledge_run(
                db,
                run,
                model=DirectFakeModel(
                    [_direct_point("企业使命是企业经营目的和社会责任的规定。")],
                    file_transport=file_transport,
                ),
            )
            assert completed.status == "completed", (
                completed.error_code,
                completed.error_message,
            )
            assert completed.statistics["candidate_count"] == 1
            assert completed.statistics["usable_count"] == 1
            assert completed.statistics["filtered_count"] == 0
            assert completed.reviewer_config == {"enabled": False}
            assert db.query(StudyDocumentUnit).filter_by(document_id=document.id).count() == 0
            assert db.query(StudyStructuredBlock).filter_by(run_id=run.id).count() == 0
            stored = db.query(StudyKnowledgeUnit).filter_by(run_id=run.id).one()
            assert stored.check_meta["local_source_blocks_persisted"] is False
            assert stored.check_meta["source_transport"] == file_transport
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_doubao_page_checkpoints_commit_each_page_and_resume_failed_page(tmp_path, monkeypatch):
    monkeypatch.setattr("core.config.settings.study_upload_dir", str(tmp_path))
    _, factory, engine = _api_fixture()
    try:
        with factory() as db:
            user = User(student_id=f"checkpoint-{uuid.uuid4().hex[:8]}", name="checkpoint", nickname="checkpoint", password_hash=hash_password("pass"))
            db.add(user)
            db.flush()
            relative = f"{user.id}/source.pdf"
            path = tmp_path / relative
            path.parent.mkdir(parents=True)
            pdf = pymupdf.open()
            try:
                for page_number in range(3):
                    page = pdf.new_page()
                    page.insert_text((72, 72), f"source page {page_number + 1}")
                pdf.save(path)
            finally:
                pdf.close()
            document = StudyDocument(user_id=user.id, original_filename="source.pdf", file_type="pdf", storage_path=relative, file_size=path.stat().st_size, sha256=hashlib.sha256(path.read_bytes()).hexdigest(), page_count=3, status="uploaded")
            db.add(document)
            db.commit()
            run, _ = create_or_reuse_run(db, document, pipeline_version="doubao-vision-v2")
            model = CheckpointDirectFakeModel(fail_page_two_attempts=2)
            failed = process_knowledge_run(db, run, model=model)
            assert failed.status == "failed"
            assert failed.statistics["processed_page_count"] == 1
            assert failed.statistics["failed_page_numbers"] == [2]
            checkpoints = db.query(StudyKnowledgePageCheckpoint).filter_by(run_id=run.id).order_by(StudyKnowledgePageCheckpoint.page_number).all()
            assert [(item.page_number, item.status, item.attempt_count) for item in checkpoints] == [(1, "completed", 1), (2, "failed", 2)]
            assert checkpoints[0].payload["knowledge_points"][0]["source_location"]["page"] == 1

            resumed = process_knowledge_run(db, failed, model=model)
            assert resumed.status == "completed"
            assert resumed.statistics["total_page_count"] == 3
            assert resumed.statistics["processed_page_count"] == 3
            assert resumed.statistics["failed_page_numbers"] == []
            assert model.page_calls == [1, 2, 2, 2, 3]
            assert db.query(StudyKnowledgeUnit).filter_by(run_id=run.id).count() == 3
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_api_requeues_failed_checkpointed_vision_run(monkeypatch):
    client, factory, engine = _api_fixture()
    monkeypatch.setattr("app.api.v1.study.process_knowledge_run_by_id", lambda _run_id: None)
    try:
        owner, headers = _create_user(factory, "checkpoint-resume")
        with factory() as db:
            document = StudyDocument(user_id=owner, original_filename="resume.pdf", file_type="pdf", storage_path="private/resume.pdf", file_size=1, sha256="c" * 64, page_count=3, status="uploaded")
            db.add(document)
            db.commit()
            run, _ = create_or_reuse_run(db, document, pipeline_version="doubao-vision-v2")
            run.status = "failed"
            run.statistics = {"total_page_count": 3, "processed_page_count": 1, "failed_page_numbers": [2]}
            run.error_code = "page_extraction_failed:2"
            db.commit()
            run_id = run.id
            document_id = document.id

        response = client.post(
            f"/api/v1/study/documents/{document_id}/knowledge-runs",
            headers=headers,
            json={"pipeline_version": "doubao-vision-v2"},
        )
        assert response.status_code == 202, response.text
        assert response.json()["data"]["reused"] is True
        assert response.json()["message"] == "已从失败页面继续处理"
        with factory() as db:
            resumed = db.get(StudyKnowledgeRun, run_id)
            assert resumed.status == "queued" and resumed.error_code is None
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_doubao_files_responses_payload_and_remote_cleanup(tmp_path, monkeypatch):
    monkeypatch.setattr("core.config.settings.study_doubao_api_key", "test-key")
    monkeypatch.setattr("core.config.settings.study_doubao_base_url", "https://ark.example/v3")
    monkeypatch.setattr("core.config.settings.study_doubao_inline_max_bytes", 0)
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append((request.method, request.url.path, request.content))
        if request.method == "POST" and request.url.path == "/v3/files":
            return httpx.Response(200, json={"id": "file-test", "status": "completed"})
        if request.method == "POST" and request.url.path == "/v3/responses":
            body = json.loads(request.content)
            assert body["model"] == "doubao-seed-2-1-turbo-260628"
            content = body["input"][0]["content"]
            assert content[0] == {"type": "input_file", "file_id": "file-test"}
            assert content[1]["type"] == "input_text"
            return httpx.Response(200, json={
                "model": body["model"],
                "output": [{
                    "type": "message",
                    "content": [{
                        "type": "output_text",
                        "text": json.dumps({"knowledge_points": []}),
                    }],
                }],
            })
        if request.method == "DELETE" and request.url.path == "/v3/files/file-test":
            return httpx.Response(200, json={"deleted": True})
        return httpx.Response(404)

    path = tmp_path / "source.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        model = DoubaoFilesResponsesModel(client=client)
        result = model.extract_file(
            path,
            filename="source.pdf",
            mime_type="application/pdf",
        )

    assert result.payload == {"knowledge_points": []}
    assert result.remote_file_deleted is True
    assert [(method, path) for method, path, _ in seen] == [
        ("POST", "/v3/files"),
        ("POST", "/v3/responses"),
        ("DELETE", "/v3/files/file-test"),
    ]


def test_doubao_small_file_uses_inline_base64_without_files_api(tmp_path, monkeypatch):
    monkeypatch.setattr("core.config.settings.study_doubao_api_key", "test-key")
    monkeypatch.setattr("core.config.settings.study_doubao_base_url", "https://ark.example/v3")
    monkeypatch.setattr("core.config.settings.study_doubao_inline_max_bytes", 1024)
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append((request.method, request.url.path))
        assert request.url.path == "/v3/responses"
        body = json.loads(request.content)
        file_input = body["input"][0]["content"][0]
        assert file_input["type"] == "input_file"
        assert file_input["filename"] == "source.pdf"
        assert file_input["file_data"].startswith("data:application/pdf;base64,")
        return httpx.Response(200, json={
            "model": body["model"],
            "output": [{
                "type": "message",
                "content": [{
                    "type": "output_text",
                    "text": json.dumps({"knowledge_points": []}),
                }],
            }],
        })

    path = tmp_path / "source.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = DoubaoFilesResponsesModel(client=client).extract_file(
            path,
            filename="source.pdf",
            mime_type="application/pdf",
        )

    assert result.file_transport == "inline_base64"
    assert result.remote_file_deleted is None
    assert seen == [("POST", "/v3/responses")]


def test_doubao_page_images_batch_and_preserve_source_page_numbers(tmp_path, monkeypatch):
    monkeypatch.setattr("core.config.settings.study_doubao_api_key", "test-key")
    monkeypatch.setattr("core.config.settings.study_doubao_base_url", "https://ark.example/v3")
    monkeypatch.setattr("core.config.settings.study_doubao_pages_per_batch", 2)
    monkeypatch.setattr("core.config.settings.study_doubao_render_dpi", 72)
    seen_batches = []

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/v3/chat/completions"
        body = json.loads(request.content)
        assert body["response_format"] == {"type": "json_object"}
        assert body["thinking"] == {"type": "disabled"}
        content = body["messages"][0]["content"]
        labels = [
            int(item["text"].split()[-1])
            for item in content
            if item["type"] == "text" and item["text"].startswith("页码 ")
        ]
        images = [item for item in content if item["type"] == "image_url"]
        assert len(images) == len(labels)
        assert all(
            item["image_url"]["url"].startswith("data:image/jpeg;base64,")
            for item in images
        )
        seen_batches.append(labels)
        points = [
            _direct_point(f"企业战略是第{number}页中完整描述的总体谋划。", page=number)
            for number in labels
        ]
        if labels == [1, 2]:
            points.append(_direct_point("企业战略还包含来源不明的其他原则。", page=99))
        return httpx.Response(200, json={
            "model": body["model"],
            "choices": [{
                "finish_reason": "stop",
                "message": {
                    "content": json.dumps({"knowledge_points": points}, ensure_ascii=False),
                },
            }],
        })

    path = tmp_path / "source.pdf"
    pdf = pymupdf.open()
    try:
        for number in range(1, 4):
            page = pdf.new_page()
            page.insert_text((72, 72), f"page {number}")
        pdf.save(path)
    finally:
        pdf.close()
    source = StudyDocument(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        original_filename="source.pdf",
        file_type="pdf",
        storage_path="source.pdf",
        file_size=path.stat().st_size,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        status="uploaded",
    )
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        records, audit, structured = extract_direct_document(
            source,
            path,
            DoubaoPageImagesResponsesModel(client=client),
        )

    assert seen_batches == [[1, 2], [3]]
    assert [item["source_refs"][0]["page_number"] for item in records] == [1, 2, 3]
    assert all(
        item["check_meta"]["source_pages_sent_as_images"] is True
        for item in records
    )
    assert all(
        item["check_meta"]["source_file_sent_directly"] is False
        for item in records
    )
    assert audit["pipeline_version"] == "doubao-vision-v2"
    assert audit["file_transport"] == "page_images"
    assert audit["model_call_count"] == 2
    assert audit["source_page_count"] == 3
    assert audit["filtered_count"] == 1
    assert audit["filtered_by_reason"]["model_risk:source_unclear"] == 1
    assert structured.parser_name == "doubao-page-image-understanding"


def test_doubao_filter_rejects_out_of_range_and_broken_delimiters():
    point = DirectKnowledgePoint.model_validate(
        _direct_point("企业使命包括（经营目的和社会责任。", page=3)
    )
    reasons = _filter_reasons(
        point,
        file_type="pdf",
        confidence_threshold=0.85,
        source_page_count=2,
    )
    assert "source_location_out_of_range" in reasons
    assert "unbalanced_delimiters" in reasons


def test_doubao_filter_rejects_course_credit_even_with_high_confidence():
    point = DirectKnowledgePoint.model_validate(_direct_point(
        "《交通系统分析》课程由某大学课程教学团队制作。",
        confidence=0.99,
    ))
    reasons = _filter_reasons(
        point,
        file_type="pdf",
        confidence_threshold=0.85,
        source_page_count=1,
    )
    assert "blocked_course_credit" in reasons


def test_doubao_structural_types_reject_false_lists_and_steps():
    false_list = DirectKnowledgePoint.model_validate(
        _direct_point("中国是一个具有悠久历史的世界文明古国。", knowledge_type="list")
    )
    false_steps = DirectKnowledgePoint.model_validate(
        _direct_point("鸦片战争揭开了社会文化转型的序幕。", knowledge_type="steps")
    )
    valid_list = DirectKnowledgePoint.model_validate(
        _direct_point("人才的基本素质包括：德；学；才；识；体。", knowledge_type="list")
    )
    valid_steps = DirectKnowledgePoint.model_validate(
        _direct_point("实施步骤：（1）确定因素；（2）确定权重。", knowledge_type="steps")
    )

    assert "list_shape_invalid" in _filter_reasons(
        false_list, file_type="pdf", confidence_threshold=0.85, source_page_count=1
    )
    assert "steps_shape_invalid" in _filter_reasons(
        false_steps, file_type="pdf", confidence_threshold=0.85, source_page_count=1
    )
    assert "list_shape_invalid" not in _filter_reasons(
        valid_list, file_type="pdf", confidence_threshold=0.85, source_page_count=1
    )
    assert "steps_shape_invalid" not in _filter_reasons(
        valid_steps, file_type="pdf", confidence_threshold=0.85, source_page_count=1
    )


def test_doubao_legacy_semantic_type_is_mapped_for_compatibility(tmp_path):
    path = tmp_path / "notes.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    source = StudyDocument(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        original_filename="notes.pdf",
        file_type="pdf",
        storage_path="notes.pdf",
        file_size=path.stat().st_size,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        status="uploaded",
    )
    legacy = _direct_point("市场细分是根据消费者需求差异划分市场。")
    legacy["type"] = "definition"

    records, _, _ = extract_direct_document(source, path, DirectFakeModel([legacy]))

    assert records[0]["knowledge_type"] == "statement"
    assert records[0]["extraction_meta"]["original_type"] == "definition"


def test_pptx_conversion_uses_temporary_pdf_and_cleans_up(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    source.write_bytes(b"pptx-test-placeholder")
    created_output = []

    def fake_run(command, *, capture_output, text, encoding, errors, timeout, check):
        assert capture_output is True and text is True and check is False
        assert encoding == "utf-8" and errors == "replace"
        assert timeout > 0
        assert command[-1] == str(source)
        output_dir = Path(command[command.index("--outdir") + 1])
        output_pdf = output_dir / "deck.pdf"
        document = pymupdf.open()
        document.new_page()
        try:
            document.save(output_pdf)
        finally:
            document.close()
        created_output.append(output_pdf)
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(
        "services.study_knowledge_doubao._resolve_soffice",
        lambda: tmp_path / "soffice-test",
    )
    monkeypatch.setattr("services.study_knowledge_doubao.subprocess.run", fake_run)
    with _source_as_pdf(source) as converted:
        assert converted.is_file()
        with pymupdf.open(converted) as pdf:
            assert pdf.page_count == 1
    assert len(created_output) == 1
    assert not created_output[0].exists()
