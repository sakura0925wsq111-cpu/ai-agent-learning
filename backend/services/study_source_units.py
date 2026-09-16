"""Deterministic, revision-scoped source units for Study knowledge V2."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER_TYPE
from pptx.oxml.ns import qn

from services.study_knowledge import (
    StructuredBlock,
    StructuredDocument,
    StructuredSourceSpan,
)


SOURCE_UNIT_VERSION = "source-units-v2"
_SENTENCE_END_RE = re.compile(r"[。！？!?]+[\"'”’）】》]*")
_TITLE_PLACEHOLDERS = {
    PP_PLACEHOLDER_TYPE.TITLE,
    PP_PLACEHOLDER_TYPE.CENTER_TITLE,
    PP_PLACEHOLDER_TYPE.SUBTITLE,
}


def _hash_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _normalized_bbox(shape, *, slide_width: int, slide_height: int) -> list[float]:
    return [
        max(0.0, min(1.0, float(shape.left) / slide_width)),
        max(0.0, min(1.0, float(shape.top) / slide_height)),
        max(0.0, min(1.0, float(shape.left + shape.width) / slide_width)),
        max(0.0, min(1.0, float(shape.top + shape.height) / slide_height)),
    ]


def _paragraph_has_bullet(paragraph) -> bool:
    properties = paragraph._p.pPr
    if properties is None:
        return False
    if properties.find(qn("a:buNone")) is not None:
        return False
    return any(
        properties.find(qn(name)) is not None
        for name in ("a:buChar", "a:buAutoNum", "a:buBlip")
    )


def _trimmed_sentence_spans(text: str) -> list[tuple[int, int]]:
    """Return non-empty sentence slices while preserving exact source offsets."""
    if not text:
        return []
    boundaries = []
    start = 0
    for match in _SENTENCE_END_RE.finditer(text):
        boundaries.append((start, match.end()))
        start = match.end()
    if start < len(text):
        boundaries.append((start, len(text)))
    if not boundaries:
        boundaries.append((0, len(text)))
    result = []
    for raw_start, raw_end in boundaries:
        segment = text[raw_start:raw_end]
        leading = len(segment) - len(segment.lstrip())
        trailing = len(segment) - len(segment.rstrip())
        item_start = raw_start + leading
        item_end = raw_end - trailing
        if item_start < item_end:
            result.append((item_start, item_end))
    return result


def _shape_is_heading(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type in _TITLE_PLACEHOLDERS
    except (AttributeError, ValueError):
        return False


def _iter_shapes(shapes, prefix: str = "") -> Iterable[tuple[str, object]]:
    """Yield leaf shapes with stable paths, including nested group shapes."""
    for shape in shapes:
        path = f"{prefix}sh{shape.shape_id}"
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes, prefix=f"{path}.")
        else:
            yield path, shape


def structured_from_pptx(path: Path, *, source_sha256: str) -> StructuredDocument:
    """Parse PPTX native objects into stable sentence/list source units."""
    presentation = Presentation(path)
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    blocks: list[StructuredBlock] = []
    reading_order = 0

    for page_number, slide in enumerate(presentation.slides, 1):
        ordered_shapes = sorted(
            _iter_shapes(slide.shapes),
            key=lambda item: (
                int(getattr(item[1], "top", 0)),
                int(getattr(item[1], "left", 0)),
                item[0],
            ),
        )
        for shape_path, shape in ordered_shapes:
            bbox = _normalized_bbox(
                shape,
                slide_width=slide_width,
                slide_height=slide_height,
            )
            base_id = f"p{page_number}:{shape_path}"
            provider_ref = f"pptx://page/{page_number}/shape/{shape_path}"

            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    cells = [cell.text for cell in row.cells]
                    rows.append("\t".join(cells))
                text = "\n".join(rows).strip()
                blocks.append(StructuredBlock(
                    id=f"{base_id}:table",
                    type="table",
                    raw_text=text,
                    reading_order=reading_order,
                    source_spans=[StructuredSourceSpan(page_number, bbox, 0, len(text))],
                    extraction_method="pptx_native",
                    quality_signals={
                        "usable_text": False,
                        "source_container_id": base_id,
                    },
                    provider_raw_ref=provider_ref,
                ))
                reading_order += 1
                continue

            if getattr(shape, "has_text_frame", False):
                heading = _shape_is_heading(shape)
                list_group_id = f"{base_id}:list"
                for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
                    paragraph_text = paragraph.text
                    if not paragraph_text.strip():
                        continue
                    paragraph_id = f"{base_id}:para{paragraph_index}"
                    is_list_item = _paragraph_has_bullet(paragraph)
                    if heading:
                        spans = [(0, len(paragraph_text))]
                        block_type = "heading"
                    elif is_list_item:
                        spans = [(0, len(paragraph_text))]
                        block_type = "list_item"
                    else:
                        spans = _trimmed_sentence_spans(paragraph_text)
                        block_type = "paragraph"

                    for sentence_index, (start, end) in enumerate(spans):
                        text = paragraph_text[start:end]
                        suffix = (
                            "heading"
                            if heading
                            else f"li{paragraph_index}"
                            if is_list_item
                            else f"s{sentence_index}"
                        )
                        blocks.append(StructuredBlock(
                            id=f"{paragraph_id}:{suffix}",
                            type=block_type,
                            raw_text=text,
                            reading_order=reading_order,
                            source_spans=[StructuredSourceSpan(
                                page_number=page_number,
                                bbox=bbox,
                                char_start=0,
                                char_end=len(text),
                            )],
                            extraction_method="pptx_native",
                            quality_signals={
                                "usable_text": True,
                                "source_container_id": paragraph_id,
                                "shape_id": int(shape.shape_id),
                                "paragraph_index": paragraph_index,
                                "shape_char_start": start,
                                "shape_char_end": end,
                                "is_heading_placeholder": heading,
                                "is_list_item": is_list_item,
                                "list_level": int(paragraph.level),
                                "list_group_id": list_group_id if is_list_item else None,
                            },
                            provider_raw_ref=provider_ref,
                        ))
                        reading_order += 1
                continue

            block_type = (
                "figure"
                if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}
                else "unknown"
            )
            blocks.append(StructuredBlock(
                id=f"{base_id}:{block_type}",
                type=block_type,
                raw_text="",
                reading_order=reading_order,
                source_spans=[StructuredSourceSpan(page_number, bbox, 0, 0)],
                extraction_method="pptx_native",
                quality_signals={
                    "usable_text": False,
                    "source_container_id": base_id,
                },
                provider_raw_ref=provider_ref,
            ))
            reading_order += 1

    revision_payload = "|".join(
        f"{block.id}:{block.type}:{_hash_text(block.raw_text)}" for block in blocks
    )
    return StructuredDocument(
        revision_id=_hash_text(
            f"{SOURCE_UNIT_VERSION}|{source_sha256}|{revision_payload}"
        ),
        source_sha256=source_sha256,
        parser_name="pptx-source-units",
        parser_version=SOURCE_UNIT_VERSION,
        parser_config={
            "sentence_split": "terminal-punctuation-with-source-offsets",
            "native_bullets": True,
            "group_shapes": True,
            "slide_width": slide_width,
            "slide_height": slide_height,
        },
        blocks=blocks,
        provider_raw_ref=str(path),
    )


def number_legacy_structured_document(document: StructuredDocument) -> StructuredDocument:
    """Split supported legacy blocks into stable sentence units without semantic grouping."""
    blocks: list[StructuredBlock] = []
    reading_order = 0
    for block in document.blocks:
        if block.type not in {"paragraph"} or not block.raw_text.strip():
            blocks.append(StructuredBlock(
                **{
                    **block.__dict__,
                    "reading_order": reading_order,
                }
            ))
            reading_order += 1
            continue
        spans = _trimmed_sentence_spans(block.raw_text)
        if len(spans) <= 1:
            blocks.append(StructuredBlock(
                **{
                    **block.__dict__,
                    "reading_order": reading_order,
                }
            ))
            reading_order += 1
            continue
        for index, (start, end) in enumerate(spans):
            text = block.raw_text[start:end]
            source_spans = [
                StructuredSourceSpan(
                    page_number=span.page_number,
                    bbox=span.bbox,
                    char_start=0,
                    char_end=len(text),
                )
                for span in block.source_spans
            ]
            signals = dict(block.quality_signals)
            signals.update({
                "source_container_id": block.id,
                "container_char_start": start,
                "container_char_end": end,
            })
            blocks.append(StructuredBlock(
                id=f"{block.id}:s{index}",
                type=block.type,
                raw_text=text,
                reading_order=reading_order,
                parent_id=block.id,
                source_spans=source_spans,
                extraction_method=block.extraction_method,
                quality_signals=signals,
                provider_raw_ref=block.provider_raw_ref,
            ))
            reading_order += 1
    revision_payload = "|".join(
        f"{block.id}:{_hash_text(block.raw_text)}" for block in blocks
    )
    return StructuredDocument(
        revision_id=_hash_text(
            f"{SOURCE_UNIT_VERSION}|{document.revision_id}|{revision_payload}"
        ),
        source_sha256=document.source_sha256,
        parser_name=f"{document.parser_name}+source-units",
        parser_version=SOURCE_UNIT_VERSION,
        parser_config={
            **document.parser_config,
            "source_unit_version": SOURCE_UNIT_VERSION,
        },
        blocks=blocks,
        provider_raw_ref=document.provider_raw_ref,
    )


__all__ = [
    "SOURCE_UNIT_VERSION",
    "number_legacy_structured_document",
    "structured_from_pptx",
]
