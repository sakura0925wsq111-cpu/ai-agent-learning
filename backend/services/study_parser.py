"""Native-text and OCR page parsing for the Study MVP."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from pptx import Presentation
from sqlalchemy.orm import Session

from core.config import settings
from models.study import StudyDocument, StudyDocumentUnit
from services.study_extraction_quality import PARAGRAPH_GATE_VERSION, assess_ocr_page
from services.study_files import LocalStudyStorage, sha256_file
from services.study_ocr import (
    OCREngine,
    OCRProcessingError,
    OCRUnavailableError,
    get_ocr_engine,
)


@dataclass(frozen=True)
class ParsedUnit:
    page_number: int
    unit_type: str
    raw_text: str
    normalized_text: str
    safe_text: str
    text_hash: str
    extraction_method: str = "native"
    ocr_confidence: float | None = None
    ocr_blocks: list[dict] = field(default_factory=list)
    quality_status: str = "accepted"
    quality_reasons: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ParsedDocument:
    page_count: int
    units: list[ParsedUnit]


class StudyParseError(Exception):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(message)


def _unit(
    page_number: int,
    unit_type: str,
    text: str,
    *,
    extraction_method: str = "native",
    ocr_confidence: float | None = None,
    ocr_blocks: list[dict] | None = None,
    safe_text: str | None = None,
    quality_status: str = "accepted",
    quality_reasons: list[str] | None = None,
) -> ParsedUnit | None:
    raw_text = text.replace("\x00", "").strip()
    normalized_text = re.sub(r"\s+", " ", raw_text).strip()
    if not normalized_text:
        return None
    return ParsedUnit(
        page_number=page_number,
        unit_type=unit_type,
        raw_text=raw_text,
        normalized_text=normalized_text,
        safe_text=(safe_text if safe_text is not None else raw_text).strip(),
        text_hash=hashlib.sha256(normalized_text.encode("utf-8")).hexdigest(),
        extraction_method=extraction_method,
        ocr_confidence=ocr_confidence,
        ocr_blocks=ocr_blocks or [],
        quality_status=quality_status,
        quality_reasons=quality_reasons or [],
    )


def _ocr_page(
    page: pymupdf.Page,
    page_number: int,
    engine: OCREngine,
) -> ParsedUnit | None:
    primary_scale = settings.study_ocr_dpi / 72
    verification_scale = settings.study_ocr_verify_dpi / 72
    primary_pixmap = page.get_pixmap(
        matrix=pymupdf.Matrix(primary_scale, primary_scale),
        colorspace=pymupdf.csRGB,
        alpha=False,
    )
    verification_pixmap = page.get_pixmap(
        matrix=pymupdf.Matrix(verification_scale, verification_scale),
        colorspace=pymupdf.csRGB,
        alpha=False,
    )
    primary_size = (primary_pixmap.width, primary_pixmap.height)
    verification_size = (verification_pixmap.width, verification_pixmap.height)
    try:
        primary_blocks = engine.recognize(primary_pixmap)
        verification_blocks = engine.recognize(verification_pixmap)
    except OCRUnavailableError as exc:
        raise StudyParseError("ocr_unavailable", str(exc)) from exc
    except OCRProcessingError as exc:
        raise StudyParseError("ocr_failed", str(exc)) from exc
    finally:
        del primary_pixmap
        del verification_pixmap
    if not primary_blocks:
        return None
    assessment = assess_ocr_page(
        primary_blocks,
        verification_blocks,
        primary_size=primary_size,
        verification_size=verification_size,
        min_confidence=settings.study_ocr_min_confidence,
    )
    return _unit(
        page_number,
        "page",
        assessment.raw_text,
        extraction_method="ocr",
        ocr_confidence=assessment.confidence,
        ocr_blocks=assessment.blocks,
        safe_text=assessment.safe_text,
        quality_status=assessment.quality_status,
        quality_reasons=assessment.quality_reasons,
    )


def _parse_pdf(
    path: Path,
    *,
    ocr_engine: OCREngine | None = None,
    page_numbers: list[int] | None = None,
) -> ParsedDocument:
    attempted_ocr = False
    try:
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise StudyParseError("password_protected_pdf", "暂不支持加密 PDF")
            page_count = len(document)
            if page_count == 0:
                raise StudyParseError("empty_document", "PDF 中没有页面")
            selected = (
                range(page_count)
                if page_numbers is None
                else [number - 1 for number in sorted(set(page_numbers))]
            )
            if any(index < 0 or index >= page_count for index in selected):
                raise StudyParseError("invalid_page_range", "OCR 页码超出 PDF 范围")

            units = []
            engine = ocr_engine
            for index in selected:
                page = document[index]
                parsed = _unit(index + 1, "page", page.get_text("text", sort=True))
                if parsed is None and (settings.study_ocr_enabled or engine is not None):
                    attempted_ocr = True
                    if engine is None:
                        try:
                            engine = get_ocr_engine()
                        except OCRUnavailableError as exc:
                            raise StudyParseError("ocr_unavailable", str(exc)) from exc
                    parsed = _ocr_page(page, index + 1, engine)
                if parsed is not None:
                    units.append(parsed)
    except StudyParseError:
        raise
    except Exception as exc:
        raise StudyParseError("invalid_pdf", "PDF 解析失败") from exc

    if not units:
        message = (
            "OCR 后仍未识别到可用文字"
            if attempted_ocr
            else "PDF 中没有可提取文字，且 OCR 未启用"
        )
        raise StudyParseError("no_extractable_text", message)
    return ParsedDocument(page_count=page_count, units=units)


def _shape_text(shape) -> str:
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append("\t".join(cells))
        return "\n".join(rows)
    if getattr(shape, "has_text_frame", False):
        return shape.text
    return ""


def _parse_pptx(path: Path) -> ParsedDocument:
    try:
        presentation = Presentation(path)
        page_count = len(presentation.slides)
        units = []
        for index, slide in enumerate(presentation.slides):
            shapes = sorted(
                slide.shapes,
                key=lambda shape: (getattr(shape, "top", 0), getattr(shape, "left", 0)),
            )
            parts = [text for shape in shapes if (text := _shape_text(shape).strip())]
            parsed = _unit(index + 1, "slide", "\n".join(parts))
            if parsed is not None:
                units.append(parsed)
    except Exception as exc:
        raise StudyParseError("invalid_pptx", "PPTX 解析失败") from exc

    if page_count == 0:
        raise StudyParseError("empty_document", "PPTX 中没有幻灯片")
    if not units:
        raise StudyParseError(
            "no_extractable_text",
            "PPTX 中没有可提取文字，暂不支持纯图片幻灯片",
        )
    return ParsedDocument(page_count=page_count, units=units)


def parse_study_file(
    file_type: str,
    path: Path,
    *,
    ocr_engine: OCREngine | None = None,
    page_numbers: list[int] | None = None,
) -> ParsedDocument:
    if file_type == "pdf":
        return _parse_pdf(path, ocr_engine=ocr_engine, page_numbers=page_numbers)
    if file_type == "pptx":
        if page_numbers is not None:
            raise StudyParseError("invalid_page_range", "PPTX 暂不支持指定解析页码")
        return _parse_pptx(path)
    raise StudyParseError("unsupported_file_type", "仅支持 PDF 或 PPTX 文件")


def _mark_failed(
    db: Session,
    document_id: str,
    error: StudyParseError,
) -> None:
    db.rollback()
    document = db.get(StudyDocument, document_id)
    if document is not None:
        document.status = "failed"
        document.error_code = error.code
        document.error_message = error.message
        db.commit()


def parse_and_store_document(
    db: Session,
    document: StudyDocument,
    storage: LocalStudyStorage | None = None,
    *,
    ocr_engine: OCREngine | None = None,
) -> int:
    """Parse one owned document and atomically replace its page/slide units."""
    existing_units = (
        db.query(StudyDocumentUnit)
        .filter(StudyDocumentUnit.document_id == document.id)
        .all()
    )
    if document.status == "parsed" and existing_units and all(
        unit.extraction_method != "ocr" or (
            bool(unit.ocr_blocks) and all(
                block.get("paragraph_gate_version") == PARAGRAPH_GATE_VERSION
                for block in unit.ocr_blocks
            )
        )
        for unit in existing_units
    ):
        return len(existing_units)

    document_id = document.id
    document.status = "parsing"
    document.error_code = None
    document.error_message = None
    db.commit()

    storage = storage or LocalStudyStorage()
    try:
        source_path = storage.resolve(document.storage_path)
        if not source_path.is_file():
            raise StudyParseError("source_file_missing", "学习资料源文件不存在")
        if sha256_file(source_path) != document.sha256:
            raise StudyParseError("source_file_changed", "学习资料源文件校验失败")
        parsed = parse_study_file(
            document.file_type,
            source_path,
            ocr_engine=ocr_engine,
        )
    except StudyParseError as exc:
        _mark_failed(db, document_id, exc)
        raise
    except OSError as exc:
        error = StudyParseError("source_file_unreadable", "学习资料源文件无法读取")
        _mark_failed(db, document_id, error)
        raise error from exc

    try:
        db.query(StudyDocumentUnit).filter(
            StudyDocumentUnit.document_id == document_id
        ).delete(synchronize_session=False)
        for unit in parsed.units:
            db.add(StudyDocumentUnit(
                document_id=document_id,
                page_number=unit.page_number,
                unit_index=0,
                unit_type=unit.unit_type,
                raw_text=unit.raw_text,
                normalized_text=unit.normalized_text,
                safe_text=unit.safe_text,
                text_hash=unit.text_hash,
                extraction_method=unit.extraction_method,
                ocr_confidence=unit.ocr_confidence,
                ocr_blocks=unit.ocr_blocks,
                quality_status=unit.quality_status,
                quality_reasons=unit.quality_reasons,
            ))
        document = db.get(StudyDocument, document_id)
        if document is None:
            raise RuntimeError("study document disappeared during parsing")
        document.page_count = parsed.page_count
        document.status = "parsed"
        document.error_code = None
        document.error_message = None
        db.commit()
    except Exception:
        db.rollback()
        raise
    return len(parsed.units)


__all__ = [
    "ParsedDocument",
    "ParsedUnit",
    "StudyParseError",
    "parse_and_store_document",
    "parse_study_file",
]
